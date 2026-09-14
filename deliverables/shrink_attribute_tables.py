#!/usr/bin/env python3
"""Set a smaller font in the two nine-column attribute tables of the L24 deliverable .docx.

  python3 deliverables/shrink_attribute_tables.py <deliverable.docx>
Rewrites the .docx IN PLACE (argv[1]); prints a [TABLE-FONT] check line, exits nonzero on a miss.

Runs AFTER pandoc in build_l24_deliverable_doc.sh. pandoc's markdown->docx has no per-table font
control, so this post-pass sets every run in each table with NINE columns (Tables 1 and 2 -- the
calibration-data and RMSE tables have 3 and 6) to PT. Text is untouched, so the docx->markdown sync
round-trip is unaffected. Written with python-docx (never raw XML: `docx_ns_prefixes`), and the
result is re-read independently below.
"""
import sys
from docx import Document
from docx.shared import Pt

PATH = sys.argv[1]
PT = 8
NCOLS = 9

d = Document(PATH)
done = 0
for t in d.tables:
    if len(t.columns) != NCOLS:
        continue
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(PT)
    done += 1
d.save(PATH)
chk = Document(PATH)
sizes = {r.font.size.pt for t in chk.tables if len(t.columns) == NCOLS
         for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.font.size}
print("[TABLE-FONT] %d nine-column table(s) set to %d pt; re-read sizes: %s" % (done, PT, sorted(sizes)))
if done != 2 or sizes != {float(PT)}:
    sys.exit("[TABLE-FONT] expected 2 tables at %d pt" % PT)
